# -*- coding: utf-8 -*-
"""生成《Roschelle 1992 阅读笔记——收敛性概念变化与实验映射》Word 文档。"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches

from build_report_docx import (
    style_font,
    run_font,
    add_numbering_defs,
    add_list_para,
    make_table,
    add_heading,
    add_body,
    add_bottom_rule,
    add_callout,
    add_page_number,
)

INK_BLUE = RGBColor(0x0B, 0x25, 0x45)
MUTED = RGBColor(0x59, 0x59, 0x59)


def add_h2(doc, text):
    return add_heading(doc, text, 2)


def build():
    doc = Document()
    doc.core_properties.title = "Roschelle 1992 阅读笔记"

    normal = doc.styles["Normal"]
    style_font(normal, size=11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in (
        ("Heading 1", 16, RGBColor(0x2E, 0x74, 0xB5), 16, 8),
        ("Heading 2", 13, RGBColor(0x2E, 0x74, 0xB5), 12, 6),
        ("Heading 3", 12, RGBColor(0x1F, 0x4D, 0x78), 8, 4),
    ):
        st = doc.styles[name]
        style_font(st, size=size, bold=True, color=color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.1

    add_numbering_defs(doc)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    fp = section.footer.paragraphs[0]
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    add_page_number(fp)

    # ---------- 首屏 ----------
    first_p = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    first_p.paragraph_format.space_before = Pt(0)
    first_p.paragraph_format.space_after = Pt(2)
    r = first_p.add_run("Roschelle (1992) 阅读笔记")
    run_font(r, size=22, bold=True, color=INK_BLUE)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after = Pt(10)
    rs = sub.add_run("Learning by Collaborating: Convergent Conceptual Change · 对实验有用的部分详解")
    run_font(rs, size=13, color=MUTED)

    for label, value in (
        ("文献", "Roschelle, J. (1992). Journal of the Learning Sciences, 2(3), 235–276."),
        ("本地文件", "06_引导形成共识_教育理论/roschelle1992.pdf（作者自存版，58 页）"),
        ("整理日期", "2026-08-31"),
        ("用途", "支撑教师模型：终止条件（论证级收敛）、测试案例、逐步提高证据标准"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.1
        rl = p.add_run(label + "：")
        run_font(rl, size=11, bold=True, color=INK_BLUE)
        rv = p.add_run(value)
        run_font(rv, size=11)
    add_bottom_rule(doc, after=12)

    # ---------- 结论速览 ----------
    add_callout(
        doc,
        "结论速览",
        "这篇论文回答了“讨论/协作凭什么能形成共识”：共识不是少数服从多数，也不是强者教弱者，"
        "而是参与者通过“展示—确认—修复”的对话循环，逐步把彼此的意义收敛到同一套深层结构上。"
        "四个过程特征（深层特征情境、隐喻交织、展示-确认-修复循环、逐步提高的证据标准）"
        "可以逐条映射到我们的教师协议：终止条件、测试案例、L1–L4 反馈、transactive 质询。",
    )
    add_body(doc, "", after=2)

    # ---------- 一、核心问题 ----------
    add_heading(doc, "一、论文在回答什么问题", 1)
    add_body(doc, "一句话：两个人（或更多人）如何在各自理解天然趋向发散的情况下，"
                  "共同建构出“共享的意义”？Roschelle 把这称为收敛问题（the problem of convergence），"
                  "并主张协作学习（learning by collaborating）的核心就是收敛（p.3 原文："
                  "“the crux of collaboration is the problem of convergence”）。")
    add_body(doc, "为什么这对科学教育尤其重要：学生带着与科学概念不兼容的前概念（misconceptions）进入课堂，"
                  "意义天然发散；而科学本身是科学家协作收敛的产物。所以一篇严肃的科学学习理论必须解释"
                  "“在发散压力下，收敛是如何达成的”，而不是假设共识会自动出现。")
    add_body(doc, "对我们的直接含义：基线三方互评正是“没有解释收敛机制”的讨论——三个模型各自发散，"
                  "最后靠平均池化强行合并。Roschelle 提醒我们：共识必须由过程产生，不能由统计硬凑。", after=2)

    # ---------- 二、三种协作理论对比 ----------
    add_heading(doc, "二、三种协作理论视角的对比（p.41–42）", 1)
    make_table(
        doc,
        [1.35, 2.2, 2.0],
        ["视角", "核心主张", "局限（Roschelle 的批评）"],
        [
            ("Vygotsky 派", "脚手架：更有能力的同伴帮弱者达到更高水平", "角色不对称，偏向“复制已有知识”"),
            ("Piaget 派", "认知冲突：失衡驱动概念变化", "只强调冲突收益，忽略共同建构"),
            ("Roschelle 收敛派", "共同意义的互相建构（mutual construction）", "论文的主张：前两者都偏离了协作的共享本质"),
        ],
        first_col_bold=True,
    )
    add_body(doc, "原文关键句（p.42）：“mutual construction of knowledge was a more apparent feature of "
                  "students' interaction than was cognitive conflict or scaffolding.”——协作中最明显的不是冲突，"
                  "也不是脚手架，而是共同建构。", after=2)
    add_body(doc, "对我们的含义：教师模型既不能只当“裁判”（不参与收敛），也不能退化成“灌输者”（直接给答案）。"
                  "它应该组织共同建构：让三个模型把彼此的意义修到同一套深层结构上。", after=2)

    # ---------- 三、四个过程特征 ----------
    add_heading(doc, "三、收敛的四个过程特征（本文最核心、对实验最有用）", 1)
    add_body(doc, "论文提出收敛性概念变化由四个特征构成（p.5、p.34）：前两个描述“概念变化往哪走”，"
                  "后两个描述“收敛的机制怎么运转”。")

    add_h2(doc, "3.1 (a) 深层特征情境（deep-featured situation）")
    add_body(doc, "参与者共同锚定一个“中间抽象层”的共享表征（论文案例里是“矢量加法三角形”），"
                  "而不是纠缠于表面特征。表面变了（屏幕显示、手势、手指、手臂），深层特征始终一致。")
    add_list_para(doc, "→ 实验映射：分歧识别和教师提问必须锚定“可检验的深层对象”（分子式碳数、反应类型、题干条件），"
                       "而不是表面的措辞差异。bad1 里三个模型的表面措辞都在说“碳数”，但深层锚点只有一个："
                       "“起始物到底几个碳、每步变不变”——教师应把讨论钉在这个深层对象上。", 110)

    add_h2(doc, "3.2 (b) 隐喻的交织（interplay of metaphors）")
    add_body(doc, "没有一个隐喻单独成立（pulling 太模糊），多个隐喻（pulling + adding + travelling + hinging）"
                  "互相约束，共同逼近概念。单独看每个都是错的，合起来却是对的近似。")
    add_list_para(doc, "→ 实验映射：对应 TAP 多要素互相约束——data（碳数）、warrant（守恒规则）、"
                       "backing（适用条件）必须一起检查，单独一个理由不可信。bad1 的 biomedical 链"
                       "“10 个碳 + 去/加碳 = 11”就是典型：单个事实错、单个步骤错，但多个错误恰好抵消。"
                       "教师应要求“每条主张至少两条独立依据互相印证”，而不是采信单个理由。", 110)

    add_h2(doc, "3.3 (c) 展示—确认—修复循环（displaying, confirming, repairing）")
    add_body(doc, "收敛靠对话轮次循环实现：一方展示自己的理解 → 对方确认 → 发现分歧就进入修复。"
                  "论文案例里，Dana 检测到明显分歧（Episode 4）时，双方立刻进入修复循环，而不是各说各话。")
    add_list_para(doc, "→ 实验映射：这就是 transactive 发言协议 + “改口必须说明采纳了哪条反馈”的理论依据。"
                       "教师每轮应显式输出状态：已确认（confirm）/ 需修复（repair）。"
                       "bad1 的 qwen 链（算出 8 却选 10）就是典型的“显示与自身不一致”，"
                       "教师在修复循环里应点名要求它“重新显示”，而不是让它继续解释。", 110)

    add_h2(doc, "3.4 (d) 逐步提高的证据标准（progressively higher standards of evidence）")
    add_body(doc, "收敛不是“看起来同意了”就算数。论文 p.36 讲得很清楚：Episode 1 两人补全了一个互相满意的解释，"
                  "但这不构成共享理解的充分证据；于是 Episode 2 他们主动构造了一个新测试案例来检验各自的预测。"
                  "证据标准随时间逐步提高：表面一致 → 测试案例 → 抽象情境中的精确预测。")
    add_list_para(doc, "→ 实验映射：这是“论证级收敛”终止条件的直接理论出处——答案字母一致只是低标准证据，"
                       "必须让模型过“测试案例”这一关：对新变体/反例做预测（例如“如果换这个试剂/条件，结论变吗”），"
                       "预测一致才算收敛。这也对应我们 L1–L4 反馈：L1 提示注意 → L2 指出矛盾 → "
                       "L3 引导重算 → L4 参考锚，本质就是证据标准逐级提高。", 110)

    # ---------- 四、案例简介 ----------
    add_heading(doc, "四、案例简介（Carol 与 Dana）", 1)
    add_body(doc, "两名学生在“Envisioning Machine”速度/加速度模拟器上协作，最终共同建构出"
                  "“矢量加法三角形”这一共享深层结构，完成了从直觉到牛顿式定性理解的收敛性概念变化。"
                  "过程跨越 5 个片段：pulling 隐喻 → 加入 hinging 约束 → 引入 adding（速度变化）→ "
                  "检测到分歧并修复 → 引入 travelling（沿直线的拉动）。")
    add_body(doc, "对我们的方法论文本价值：这是一篇“单案例微观分析”，Roschelle 自己承认案例研究不能证明理论，"
                  "但能澄清概念并吸引注意。汇报时不要把它说成“实验证据”，而应定位为“机制理论 + 过程框架”，"
                  "实验证据由 Weinberger 2010 提供。", after=2)

    # ---------- 五、实验映射汇总 ----------
    add_heading(doc, "五、对实验的映射汇总（可直接写进方案）", 1)
    make_table(
        doc,
        [1.7, 2.3, 1.7],
        ["Roschelle 概念", "我们协议中的对应部件", "具体改动"],
        [
            ("深层特征情境", "分歧识别 + 提问锚点", "分歧清单按“事实层→推理层→结构层”组织，钉住可检验对象"),
            ("隐喻交织", "TAP 多要素约束", "要求每条主张至少两条独立依据互相印证"),
            ("展示-确认-修复循环", "transactive 质询 + 改口规则", "教师每轮显式输出 confirm/repair 状态"),
            ("逐步提高证据标准", "L1–L4 反馈 + 终止条件", "终态从“答案一致”升级为“测试案例预测一致”"),
            ("共同建构（非灌输/非冲突）", "教师角色定位", "教师只组织不答题；目标是共享深层结构而非输赢"),
        ],
        first_col_bold=True,
    )

    # ---------- 六、可直接落地的三条改动 ----------
    add_heading(doc, "六、三条可直接落地的设计改动（建议）", 1)
    add_list_para(doc, "改动 1：终止条件加“测试案例”一级。模型不仅需答案一致、理由相容，还需对教师给出的"
                       "一个反例/变体做预测且预测一致；预测冲突则继续一轮（对应 3.4，成本可控：只对已收敛题抽查）。", 110)
    add_list_para(doc, "改动 2：教师每轮输出“confirm/repair”状态标签。发言若只是复述原推理（非 transactive）"
                       "标 repair；若确实加工了他人推理且补上了证据，标 confirm。这让“是否在收敛”变成可观测的过程指标。", 110)
    add_list_para(doc, "改动 3：提问锚定深层对象。分歧清单里每条冲突必须落到“一个可检验的深层特征”（数字、"
                       "试剂、条件），禁止停留在“你的说法和我的说法不同”这种表面描述。", 110)

    # ---------- 七、可引用原句 ----------
    add_heading(doc, "七、可引用原句（附页码）", 1)
    for q in [
        "“the crux of collaboration is the problem of convergence” (p.3)",
        "“convergent conceptual change is achieved incrementally, interactively, and socially through collaborative participation in joint activity” (p.5)",
        "“convergence is achieved through cycles of displaying, confirming and repairing shared meanings” (p.5)",
        "“students follow confirmation of shared meaning at low standards of evidence with confirmation at high enough standards that the attribution of shared knowledge is warranted” (p.10)",
        "“although the students collaboratively completed a mutually satisfactory explanation in Episode 1, this was not sufficient evidence that they shared an understanding. Thus, it was entirely appropriate that in Episode 2, the students collaboratively produced a new test case” (p.36)",
        "“mutual construction of knowledge was a more apparent feature of students' interaction than was cognitive conflict or scaffolding” (p.42)",
    ]:
        add_list_para(doc, q, 110)

    add_body(doc, "引用建议：汇报或论文中，把 Roschelle 定位为“共识 = 收敛过程”的机制理论（非实验证据），"
                  "把 Weinberger 2010 定位为“脚本化引导有效”的实验证据，两者互补。", after=4)

    out = "/Users/Zhuanz/Documents/Codex/2026-08-02/xia/outputs/06_引导形成共识_教育理论/Roschelle1992_阅读笔记.docx"
    doc.save(out)
    print("saved:", out)


if __name__ == "__main__":
    build()
