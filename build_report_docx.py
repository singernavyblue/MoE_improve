# -*- coding: utf-8 -*-
"""生成《评审会意见整理与后续研究规划》汇报文档。"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ---------- 设计 tokens（standard_business_brief） ----------
LATIN = "Calibri"
EAST = "PingFang SC"
BODY_SIZE = 11
TABLE_SIZE = 10.5
HEADING_BLUE = RGBColor(0x2E, 0x74, 0xB5)
HEADING_DARK = RGBColor(0x1F, 0x4D, 0x78)
INK_BLUE = RGBColor(0x0B, 0x25, 0x45)
MUTED = RGBColor(0x59, 0x59, 0x59)
TABLE_HEADER_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
BORDER_GRAY = "BFBFBF"


def style_font(style, latin=LATIN, east=EAST, size=None, bold=None, color=None):
    """设置样式字体（含中文字体）。"""
    style.font.name = latin
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


def run_font(run, size=None, bold=None, color=None, italic=None):
    """设置 run 字体（含中文字体）。"""
    run.font.name = LATIN
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), LATIN)
    rfonts.set(qn("w:hAnsi"), LATIN)
    rfonts.set(qn("w:eastAsia"), EAST)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic


def add_numbering_defs(doc):
    """添加真实的 bullet / decimal 编号定义。"""
    numbering_el = doc.part.numbering_part.numbering_definitions._numbering
    for abstract_id, fmt, txt, num_id in (
        (100, "bullet", "•", 110),
        (101, "decimal", "%1.", 120),
    ):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), txt)
        lvl.append(lvl_text)
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        lvl.append(lvl_jc)
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        ppr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        ppr.append(spacing)
        lvl.append(ppr)
        rpr = OxmlElement("w:rPr")
        rfonts = OxmlElement("w:rFonts")
        rfonts.set(qn("w:ascii"), LATIN)
        rfonts.set(qn("w:hAnsi"), LATIN)
        rfonts.set(qn("w:eastAsia"), EAST)
        rpr.append(rfonts)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "22")
        rpr.append(sz)
        lvl.append(rpr)
        abstract.append(lvl)
        numbering_el.append(abstract)
        if fmt == "decimal":
            # 两组独立编号列表：问题（120）与推进策略（121），各自从 1 开始
            for extra_num_id in (num_id, num_id + 1):
                num = OxmlElement("w:num")
                num.set(qn("w:numId"), str(extra_num_id))
                lvl_override = OxmlElement("w:lvlOverride")
                lvl_override.set(qn("w:ilvl"), "0")
                start_override = OxmlElement("w:startOverride")
                start_override.set(qn("w:val"), "1")
                lvl_override.append(start_override)
                num.append(lvl_override)
                ref = OxmlElement("w:abstractNumId")
                ref.set(qn("w:val"), str(abstract_id))
                num.append(ref)
                numbering_el.append(num)
        else:
            num = OxmlElement("w:num")
            num.set(qn("w:numId"), str(num_id))
            ref = OxmlElement("w:abstractNumId")
            ref.set(qn("w:val"), str(abstract_id))
            num.append(ref)
            numbering_el.append(num)


def add_list_para(doc, text, num_id, lead=None):
    """带真实编号的列表段落，支持加粗引导词。"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    ppr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    ppr.insert_element_before(num_pr, "w:spacing", "w:ind", "w:jc")
    if lead:
        r1 = p.add_run(lead)
        run_font(r1, bold=True)
    r2 = p.add_run(text)
    run_font(r2)
    return p


def set_table_geometry(table, widths_in):
    """显式表格几何：tblW / tblInd / tblGrid / tcW 一致。"""
    widths_dxa = [int(w * 1440) for w in widths_in]
    total_dxa = 9360
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert_element_before(tbl_w, "w:jc", "w:tblInd", "w:tblBorders", "w:shd")
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.insert_element_before(tbl_ind, "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    for gc in grid.findall(qn("w:gridCol")):
        grid.remove(gc)
    for w in widths_dxa:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER_GRAY, sz="4"):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.insert_element_before(tbl_borders, "w:shd", "w:tblLayout", "w:tblCellMar")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tbl_borders.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge)
            tbl_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_cell_margins(table, top=80, bottom=80, start=120, end=120):
    tbl_pr = table._tbl.tblPr
    mar = tbl_pr.find(qn("w:tblCellMar"))
    if mar is None:
        mar = OxmlElement("w:tblCellMar")
        tbl_pr.insert_element_before(mar, "w:tblLook")
    for name, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = mar.find(qn("w:" + name))
        if el is None:
            el = OxmlElement("w:" + name)
            mar.append(el)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def repeat_header(row):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def fill_cell(cell, text, bold=False, size=TABLE_SIZE, color=None, align=None):
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    run_font(r, size=size, bold=bold, color=color)


def make_table(doc, widths_in, header, rows, header_fill=TABLE_HEADER_FILL, col_aligns=None, first_col_bold=False):
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, widths_in)
    set_table_borders(table)
    set_cell_margins(table)
    for j, h in enumerate(header):
        fill_cell(table.rows[0].cells[j], h, bold=True, color=INK_BLUE)
        shade_cell(table.rows[0].cells[j], header_fill)
    repeat_header(table.rows[0])
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            align = col_aligns[j] if col_aligns else None
            fill_cell(
                table.rows[i + 1].cells[j],
                val,
                bold=(first_col_bold and j == 0),
                align=align,
            )
    return table


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        run_font(r, bold=True)
    return h


def add_body(doc, text, after=6, size=BODY_SIZE, color=None, bold=False, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_rich(doc, parts, after=6):
    """parts: list of (text, dict) 依次加入 run。"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    for text, kwargs in parts:
        r = p.add_run(text)
        run_font(r, **kwargs)
    return p


def add_bottom_rule(doc, color="2E74B5", sz="12", space="6", after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    ppr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    ppr.append(p_bdr)
    return p


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [6.5])
    set_cell_margins(table, top=120, bottom=120, start=160, end=160)
    shade_cell(table.rows[0].cells[0], CALLOUT_FILL)
    p = table.rows[0].cells[0].paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r1 = p.add_run(label + "：")
    run_font(r1, size=11, bold=True, color=HEADING_DARK)
    r2 = p.add_run(text)
    run_font(r2, size=11)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    run_font(run, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._element.append(fld_begin)
    run._element.append(instr)
    run._element.append(fld_end)


def build():
    doc = Document()
    doc.core_properties.title = "评审会意见整理与后续研究规划"

    # ---------- 样式 ----------
    normal = doc.styles["Normal"]
    style_font(normal, size=11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in (
        ("Heading 1", 16, HEADING_BLUE, 16, 8),
        ("Heading 2", 13, HEADING_BLUE, 12, 6),
        ("Heading 3", 12, HEADING_DARK, 8, 4),
    ):
        st = doc.styles[name]
        style_font(st, size=size, bold=True, color=color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.1

    add_numbering_defs(doc)

    # ---------- 页面与页脚 ----------
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    add_page_number(fp)

    # ---------- 首屏（memo_masthead） ----------
    first_p = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    first_p.paragraph_format.space_before = Pt(0)
    first_p.paragraph_format.space_after = Pt(2)
    r = first_p.add_run("评审会意见整理与后续研究规划")
    run_font(r, size=22, bold=True, color=INK_BLUE)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after = Pt(10)
    rs = sub.add_run("第三方裁判升级方向——从“静态打分”到“教师模型”引导式评估")
    run_font(rs, size=13, color=MUTED)

    for label, value in (
        ("汇报人", "王行杭"),
        ("日期", "2026-08-28"),
        ("会议主题", "评审会 · 第三方裁判 / 能力画像方向"),
        ("状态", "待执行（两步推进）"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.1
        rl = p.add_run(label + "：")
        run_font(rl, size=11, bold=True, color=INK_BLUE)
        rv = p.add_run(value)
        run_font(rv, size=11)

    add_bottom_rule(doc, after=12)

    # ---------- 结论速览 ----------
    add_callout(
        doc,
        "会议结论速览",
        "当前实验方向有效但缺少理论闭环。下一步核心是：① 用教育学理论（建构主义、对比学习）补齐评估标准的理论依据；"
        "② 将第三方从“静态打分裁判”升级为“会提问的教师模型”，通过师生式问答引导模型暴露盲区、形成共识。",
    )
    add_body(doc, "", after=2)

    # ---------- 一、当前基线 ----------
    add_heading(doc, "一、当前基线（会议确认）", 1)
    add_list_para(
        doc,
        "第三方 Prometheus 仅作为“裁判”对最终结果做综合判断，不参与内容生成；此前“用第三方模型指导原始图传播权重”的路线已确认放弃。",
        110,
        lead="方案定位：",
    )
    add_list_para(
        doc,
        "引入第三方后，较原版模型互评基线提升约 6 个百分点。",
        110,
        lead="实验效果：",
    )
    add_list_para(
        doc,
        "一维标准 → 四维标准（7B 上效果略有下降）→ DeepSeek 大模型（选链与正确数一致，排除模型能力干扰）→ 单维绝对打分（较四维提升约 5 个百分点）。结论：问题不在模型大小，而在评估标准本身的逻辑。",
        110,
        lead="评估标准演进：",
    )
    add_list_para(
        doc,
        "引入更大量级推理模型未带来提升，需确认其是具备特定评价对象的专有知识，还是仅为通用模型。",
        110,
        lead="待排查问题：",
    )

    # ---------- 二、核心问题 ----------
    add_heading(doc, "二、导师指出的两个核心问题", 1)
    add_list_para(
        doc,
        "判断逻辑缺乏理论支撑。当前规则是“倾向于选择更完整且无错误的推理链”，停留在“看起来对”的直观层面；需要溯源论文与教育学理论，形成严密的论证闭环。",
        120,
        lead="问题一：",
    )
    add_list_para(
        doc,
        "评估机制是“闭卷考试”。各模型独立作答后由第三方直接累加打分，缺乏动态交互，无法暴露和修正认知盲区。",
        120,
        lead="问题二：",
    )

    # ---------- 三、理论方向 ----------
    add_heading(doc, "三、导师给出的理论方向（教育学）", 1)
    add_list_para(
        doc,
        "通过讨论引导形成共识是教育学的成熟方法，可用于规范多模型的共识形成机制。",
        110,
        lead="建构主义与对比学习：",
    )
    add_list_para(
        doc,
        "正式讨论前先做摸底测试，建立每个模型的能力基准线（与“能力画像”方向衔接）。",
        110,
        lead="入学测试（能力基线）：",
    )
    add_list_para(
        doc,
        "从“无组织小组讨论”升级为“有教师引导的研讨课”。",
        110,
        lead="角色类比：",
    )

    # ---------- 四、教师模型机制 ----------
    add_heading(doc, "四、新机制设计：第三方升级为“教师模型”", 1)
    add_list_para(
        doc,
        "模型间辩论出现冲突时，教师模型针对分歧点提问，引导各模型输出中间结论。",
        110,
        lead="冲突触发提问：",
    )
    add_list_para(
        doc,
        "通过反复追问促使模型形成共识或自我修正，暴露认知盲区。",
        110,
        lead="多轮收敛：",
    )
    add_list_para(
        doc,
        "教师模型在提问与听取回答的过程中，同步增强自身对该问题的判断能力。",
        110,
        lead="双向赋能：",
    )
    add_list_para(
        doc,
        "提问轮次与终止条件由我们自定义，最终拍板权保留在己方。",
        110,
        lead="控制权在己方：",
    )

    add_heading(doc, "机制升级对比", 2)
    make_table(
        doc,
        [1.5, 2.5, 2.5],
        ["维度", "现状（静态打分）", "目标（教师模型）"],
        [
            ("打分方式", "模型独立作答后，第三方直接累加打分", "冲突时动态追问，按中间结论逐步评估"),
            ("交互方式", "无动态交互（闭卷考试）", "师生式深度提问（研讨课）"),
            ("共识形成", "依赖静态池化，正确少数易被淹没", "引导式提问逼出共识或自我修正"),
            ("第三方角色", "只做裁判，不参与生成", "边听边学、双向赋能的“教师”"),
            ("理论依据", "主观直觉", "建构主义、对比学习、入学测试"),
        ],
        first_col_bold=True,
    )
    add_body(doc, "", after=4)

    # ---------- 五、两步推进 ----------
    add_heading(doc, "五、两步推进策略", 1)
    add_list_para(
        doc,
        "夯实理论依据——为现有四维评估标准找到教育学 / 测量学原始出处，形成“为什么一条推理链是好的”论证闭环。",
        121,
        lead="第一步：",
    )
    add_list_para(
        doc,
        "改造交互机制——将模型间辩论转化为师生式深度提问，用引导式问答替代平行辩论，通过反复提问推动观点收敛。",
        121,
        lead="第二步：",
    )

    # ---------- 六、任务清单 ----------
    add_heading(doc, "六、后续任务清单", 1)
    make_table(
        doc,
        [0.9, 2.2, 3.4],
        ["优先级", "任务", "核心要点"],
        [
            ("P0", "评估标准理论溯源", "为“选更完整且无错误的推理链”找到教育学与测量学的原始依据（建构主义、对比学习、能力评估）"),
            ("P0", "入学测试设计", "定义摸底题集、能力维度与建档方式，衔接“能力画像”方向"),
            ("P1", "教师模型提问协议", "明确提问时机（冲突点）、提问内容、中间结论如何进入下一轮评估"),
            ("P1", "终止条件定义", "定义轮次上限、置信度阈值等收敛条件，控制权保留在己方"),
            ("P2", "双向赋能验证", "设计实验验证教师模型“边听边学”是否成立，及对最终准确率的影响"),
        ],
        first_col_bold=True,
    )
    add_body(doc, "", after=4)

    # ---------- 七、待对齐问题 ----------
    add_heading(doc, "七、需进一步对齐的问题", 1)
    add_list_para(
        doc,
        "入学测试题集的来源与防泄漏（摸底题与评测题如何隔离）。",
        110,
    )
    add_list_para(
        doc,
        "教师模型提问的成本（额外轮次与推理开销）与收益的平衡。",
        110,
    )
    add_list_para(
        doc,
        "理论溯源的范围：只补教育学，还是同时纳入测量学（IRT / 认知诊断）。",
        110,
    )

    out = "/Users/Zhuanz/Documents/Codex/2026-08-02/xia/outputs/评审会意见整理与后续研究规划.docx"
    doc.save(out)
    print("saved:", out)


if __name__ == "__main__":
    build()
