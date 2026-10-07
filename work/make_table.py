#!/usr/bin/env python3
"""Create a PPT-ready xlsx table for the judge experiment comparison."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = "outputs/裁判实验对比表.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "对比"

# Title
ws.merge_cells("A1:C1")
ws["A1"] = "GoA 边采样裁判实验对比（GPQA Diamond 198 题）"
ws["A1"].font = Font(size=14, bold=True)
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 30

# Header
headers = ["变体", "全量 198", "严格子集（节点一致）"]
header_fill = PatternFill("solid", fgColor="4472C4")
header_font = Font(bold=True, color="FFFFFF", size=12)
for col, h in enumerate(headers, 1):
    c = ws.cell(row=2, column=col, value=h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 24

rows = [
    ("原版互评基线", "35.86%", "35.12%（168 题）"),
    ("无 gold 裁判", "32.32%", "30.95%"),
    ("有 gold 裁判", "41.92%", "39.18%（171 题）"),
]
thin = Side(style="thin", color="999999")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
highlight = PatternFill("solid", fgColor="FCE4D6")
for r, row in enumerate(rows, 3):
    for col, v in enumerate(row, 1):
        c = ws.cell(row=r, column=col, value=v)
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.font = Font(size=12)
        if r == 4:  # highlight the no-gold row
            c.fill = highlight
            c.font = Font(size=12, bold=True)
    ws.row_dimensions[r].height = 22

ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 26

wb.save(OUT)
print("saved", OUT)
