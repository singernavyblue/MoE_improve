#!/usr/bin/env python3
"""Create a PPT-ready xlsx for the three classic bad cases across three runs."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = "outputs/经典BadCase三版本对照表.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "BadCase对照"

# Title
ws.merge_cells("A1:F1")
ws["A1"] = "三个经典 Bad Case 在三版运行中的表现（GPQA）"
ws["A1"].font = Font(size=14, bold=True)
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 30

# Header
headers = ["Bad case", "类型", "gold", "基线互评", "有 gold 裁判", "无 gold 裁判"]
header_fill = PatternFill("solid", fgColor="4472C4")
header_font = Font(bold=True, color="FFFFFF", size=12)
for col, h in enumerate(headers, 1):
    c = ws.cell(row=2, column=col, value=h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 24

rows = [
    ("bad1 肉桂醛碳数", "上下文脆弱性", "C", "C ✓", "D ✗", "C ✓"),
    ("bad2 SARS except", "指令理解+回音室", "D", "C ✗", "B ✗", "B ✗"),
    ("bad3 立体化学", "谄媚从众", "B", "A ✗", "A ✗", "A ✗"),
]
thin = Side(style="thin", color="999999")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
green = PatternFill("solid", fgColor="C6EFCE")
red = PatternFill("solid", fgColor="FFC7CE")

for r, row in enumerate(rows, 3):
    for col, v in enumerate(row, 1):
        c = ws.cell(row=r, column=col, value=v)
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.font = Font(size=12)
        # color the result cells (columns D-F) by correctness
        if col >= 4:
            if "✓" in str(v):
                c.fill = green
                c.font = Font(size=12, bold=True)
            else:
                c.fill = red
    ws.row_dimensions[r].height = 22

# Footnote
ws.merge_cells("A7:F7")
ws["A7"] = "bad2、bad3 两版裁判均未改善：裁判只能选链、不能造链，指令误解与从众无法靠打分修复"
ws["A7"].font = Font(size=10, italic=True, color="808080")

ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 20
ws.column_dimensions["C"].width = 8
ws.column_dimensions["D"].width = 12
ws.column_dimensions["E"].width = 14
ws.column_dimensions["F"].width = 14

wb.save(OUT)
print("saved", OUT)
