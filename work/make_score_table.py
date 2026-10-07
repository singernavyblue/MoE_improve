#!/usr/bin/env python3
"""Create a PPT-ready xlsx for the no-gold judge score distribution."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = "outputs/裁判打分分布表.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "打分分布"

# Title
ws.merge_cells("A1:C1")
ws["A1"] = "无 gold 裁判打分分布（594 次）"
ws["A1"].font = Font(size=14, bold=True)
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 30

# Header
headers = ["分数", "次数", "占比"]
header_fill = PatternFill("solid", fgColor="4472C4")
header_font = Font(bold=True, color="FFFFFF", size=12)
for col, h in enumerate(headers, 1):
    c = ws.cell(row=2, column=col, value=h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 24

rows = [
    (1, 28, "4.7%"),
    (2, 68, "11.4%"),
    (3, 64, "10.8%"),
    (4, 354, "59.6%"),
    (5, 80, "13.5%"),
]
thin = Side(style="thin", color="999999")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
highlight = PatternFill("solid", fgColor="FCE4D6")
for r, row in enumerate(rows, 3):
    for col, v in enumerate(row, 1):
        c = ws.cell(row=r, column=col, value=v)
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.font = Font(size=12)
        if row[0] == 4:  # highlight the dominant 4-score row
            c.fill = highlight
            c.font = Font(size=12, bold=True)
    ws.row_dimensions[r].height = 22

# Footnote
ws.merge_cells("A8:C8")
ws["A8"] = "平均分 3.66 —— 近六成链被打 4 分，裁判区分度极低"
ws["A8"].font = Font(size=10, italic=True, color="808080")

ws.column_dimensions["A"].width = 12
ws.column_dimensions["B"].width = 12
ws.column_dimensions["C"].width = 12

wb.save(OUT)
print("saved", OUT)
