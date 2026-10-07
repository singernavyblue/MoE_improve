# -*- coding: utf-8 -*-
"""生成《教育理论基础文献清单——导师方向》Word 文档。"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches

from build_report_docx import (
    style_font,
    run_font,
    add_numbering_defs,
    add_list_para,
    make_table,
    add_heading,
    add_body,
    add_rich,
    add_bottom_rule,
    add_callout,
    add_page_number,
)

INK_BLUE = RGBColor(0x0B, 0x25, 0x45)
MUTED = RGBColor(0x59, 0x59, 0x59)


def add_ref(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    r = p.add_run(text)
    run_font(r, size=10)


def build():
    doc = Document()
    doc.core_properties.title = "教育理论基础文献清单"

    normal = doc.styles["Normal"]
    style_font(normal, size=11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in (
        ("Heading 1", 16, RGBColor(0x2E, 0x74, 0xB5), 16, 8),
        ("Heading 2", 13, RGBColor(0x2E, 0x74, 0xB5), 12, 6),
        ("Heading 3", 12, RGBColor(0x1F, 0x4D, 0x78), 8, 4),
    ):
        st = doc.styles[name]
        style_font(st, size=size, bold=True, color=color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.1

    add_numbering_defs(doc)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    fp = section.footer.paragraphs[0]
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    add_page_number(fp)

    # ---------- 首屏 ----------
    first_p = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    first_p.paragraph_format.space_before = Pt(0)
    first_p.paragraph_format.space_after = Pt(2)
    r = first_p.add_run("教育理论基础文献清单")
    run_font(r, size=22, bold=True, color=INK_BLUE)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after = Pt(10)
    rs = sub.add_run("对应导师方向：建构主义 · 对比学习 · 讨论引导共识 · 评估知识正确性与构建联系")
    run_font(rs, size=13, color=MUTED)

    for label, value in (
        ("汇报人", "王行杭"),
        ("日期", "2026-08-30"),
        ("任务", "教育理论与文献收集"),
        ("状态", "清单完成（文献均已核实出处）"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.1
        rl = p.add_run(label + "：")
        run_font(rl, size=11, bold=True, color=INK_BLUE)
        rv = p.add_run(value)
        run_font(rv, size=11)
    add_bottom_rule(doc, after=12)

    # ---------- 结论速览 ----------
    add_callout(
        doc,
        "结论速览",
        "导师提到的“教育学成熟理论”可归为六条主线：建构主义、社会建构主义与脚手架、讨论与共识形成、"
        "苏格拉底式提问、对比学习、评估知识正确性与构建联系。每条都有可直接引用的经典文献，"
        "且与我们的教师模型机制、Q 矩阵标定、四维评分标准一一对应。",
    )
    add_body(doc, "", after=2)

    # ---------- 一、总览 ----------
    add_heading(doc, "一、总览：六条主线 × 关键文献 × 对实验的支撑", 1)
    make_table(
        doc,
        [1.3, 1.6, 1.9, 1.7],
        ["主线", "核心理论", "关键文献（代表作）", "支撑实验的哪部分"],
        [
            ("建构主义", "知识是主动建构的，结论对≠理解对", "Piaget；von Glasersfeld（1989）", "评推理链而非只看答案"),
            ("社会建构主义与脚手架", "引导下达到独自达不到的水平", "Vygotsky（1978）；Wood, Bruner & Ross（1976）", "教师模型 L1–L4 分级反馈"),
            ("讨论与共识", "不同类型的对话质量不同", "Mercer（1996）；Osborne（2010）；Kuhn（1991）", "教师主持的研讨协议"),
            ("苏格拉底式提问", "引导式提问提升批判性思维", "van Zee & Minstrell（1997）；Paul & Elder", "七类提问模板"),
            ("对比学习", "比较案例促进学习", "Alfieri 等（2013）元分析；Chi（2009）ICAP", "引导式讨论优于平行动作"),
            ("评估知识与构建联系", "用证据推断能力、评估结构完整", "Facione（1990）；Toulmin（1958）；Biggs & Collis（1982）；Mislevy（2003）", "四维评分标准与 Q 矩阵"),
        ],
        first_col_bold=True,
    )
    add_body(doc, "", after=4)

    # ---------- 二、分主题 ----------
    add_heading(doc, "二、分主题理论要点与文献", 1)

    add_h2(doc, "2.1 建构主义（Constructivism）")
    add_body(doc, "核心：理解是学习者主动建构的，记住答案不等于理解；评估应探测深层理解而非核对结论。")
    add_list_para(doc, "Piaget（1952）The Origins of Intelligence in Children：认知发展是主动建构过程。", 110)
    add_list_para(doc, "von Glasersfeld（1989）Cognition, Construction of Knowledge, and Teaching：激进建构主义，知识是经验的组织而非客观复制。", 110)
    add_list_para(doc, "→ 支撑：为什么我们的裁判要看推理链而不是只对答案（bad1“结论对但推理错”必须扣分）。", 110)

    add_h2(doc, "2.2 社会建构主义与脚手架（Scaffolding）")
    add_body(doc, "核心：在更有能力的引导者帮助下，学习者能达到独自达不到的水平（最近发展区 ZPD）；教师通过脚手架逐步撤除支持。")
    add_list_para(doc, "Vygotsky（1978）Mind in Society：ZPD、语言中介、引导。", 110)
    add_list_para(doc, "Wood, Bruner & Ross（1976）The Role of Tutoring in Problem Solving（JCPP 17: 89–100）：“脚手架”一词的原始出处。", 110)
    add_list_para(doc, "Collins, Brown & Newman（1989）Cognitive Apprenticeship：教师方法包括示范、教练、脚手架、消退——教师模型机制的直接原型。", 110)
    add_list_para(doc, "Palincsar & Brown（1984）Reciprocal Teaching（Cognition and Instruction 1(2)）：以提问驱动的师生对话协议。", 110)
    add_list_para(doc, "→ 支撑：教师模型“只给恰好够用的线索、逐级加强、最终撤除”的 L1–L4 反馈设计。", 110)

    add_h2(doc, "2.3 讨论与共识形成（Talk & Consensus）")
    add_body(doc, "核心：不是所有“讨论”都有效——对话有不同类型，引导得当的对话才能推动共识与学习。")
    add_list_para(doc, "Mercer（1996）The Quality of Talk in Children's Collaborative Activity（Learning and Instruction 6(4)）：区分争执型、累积型、探索型对话——探索型才能达成高质量共识；直接对应“无组织小组讨论 vs 引导式研讨”。", 110)
    add_list_para(doc, "Osborne（2010）Arguing to Learn in Science（Science 328: 463–466）：科学教育中论证性讨论促进学习。", 110)
    add_list_para(doc, "Kuhn（1991）The Skills of Argument（Cambridge UP）：日常论证技能的发展与缺陷。", 110)
    add_list_para(doc, "Scardamalia & Bereiter（1994）Computer Support for Knowledge-Building Communities：知识建构社区与进步性话语。", 110)
    add_list_para(doc, "van Eemeren & Grootendorst（2004）A Systematic Theory of Argumentation：批判性讨论四阶段与十条规则。", 110)
    add_list_para(doc, "→ 支撑：教师研讨协议的四阶段流程、交叉质询、终止条件。", 110)

    add_h2(doc, "2.4 苏格拉底式提问（Socratic Questioning）")
    add_body(doc, "核心：用一连串追问引导对方暴露假设、澄清概念、检查证据，而不是直接给答案。")
    add_list_para(doc, "van Zee & Minstrell（1997）Using Questioning to Guide Student Thinking（JLS 6(2)）：教师提问如何引导学生思考。", 110)
    add_list_para(doc, "Paul & Elder（2007）Critical Thinking: Learn the Tools the Best Thinkers Use：苏格拉底式提问的问题分类（澄清、假设、证据、观点、影响等）。", 110)
    add_list_para(doc, "Graesser & Person（1994）Question Asking During Tutoring：辅导中提问的类型学与功能。", 110)
    add_list_para(doc, "→ 支撑：教师模型的七类提问模板（澄清/证据/推理/支援/限定/反驳/一致）。", 110)

    add_h2(doc, "2.5 对比学习（Learning Through Comparison）")
    add_body(doc, "核心：让学生比较案例/观点能显著提升学习；交互性学习优于单纯的信息接收。")
    add_list_para(doc, "Alfieri, Nokes-Malach & Schunn（2013）Learning Through Case Comparisons: A Meta-Analytic Review（Educational Psychologist 48(2)）：元分析证明案例对比学习的有效性。", 110)
    add_list_para(doc, "Chi（2009）Active-Constructive-Interactive（Topics in Cognitive Science 1(1)）：ICAP 框架——交互（Interactive）> 建构 > 主动 > 被动；讨论之所以有用是因为它是交互式学习。", 110)
    add_list_para(doc, "→ 支撑：为什么“让模型互相比较、交叉质询”比“各自独立作答”更可能产出更好共识；也解释了当前平行动作不够的原因。", 110)

    add_h2(doc, "2.6 评估知识正确性与构建联系")
    add_body(doc, "核心：评估知识是否正确、推理是否完整、概念是否建立联系，是教育测量学的成熟领域。")
    add_list_para(doc, "Facione（1990）Critical Thinking: A Statement of Expert Consensus（Delphi Report）：批判性思维六技能。", 110)
    add_list_para(doc, "Toulmin（1958）The Uses of Argument：论证六要素，定义“完整推理”。", 110)
    add_list_para(doc, "Biggs & Collis（1982）Evaluating the Quality of Learning（SOLO）：学习结果的结构层级。", 110)
    add_list_para(doc, "Novak & Gowin（1984）Learning How to Learn：概念图评估知识关联的完整度（“构建联系”）。", 110)
    add_list_para(doc, "Mislevy, Steinberg & Almond（2003）On the Structure of Educational Assessments：证据中心设计，从证据推断能力。", 110)
    add_list_para(doc, "Brookhart（2013）How to Create and Use Rubrics：评分量表的设计原则。", 110)
    add_list_para(doc, "→ 支撑：我们的四维评分标准（审题对齐/证据与推理/验证/结构深度）与 Q 矩阵标定指南。", 110)

    # ---------- 三、阅读顺序 ----------
    add_heading(doc, "三、推荐阅读顺序", 1)
    add_list_para(doc, "先读 Scaffolding 原始文献（Wood/Bruner/Ross 1976）与 Vygotsky（1978）——教师模型的理论源头。", 110)
    add_list_para(doc, "再读 Mercer（1996）与 Chi（2009）——理解“讨论为什么有用、什么讨论才有效”。", 110)
    add_list_para(doc, "然后读 van Zee & Minstrell（1997）与 Palincsar & Brown（1984）——提问协议怎么设计。", 110)
    add_list_para(doc, "最后读评估类（Toulmin/SOLO/Facione/Mislevy）——对应四维评分标准与 Q 矩阵（已在理论溯源报告引用）。", 110)

    # ---------- 参考文献 ----------
    add_heading(doc, "四、参考文献（按主题归类）", 1)
    for ref in [
        "Vygotsky, L. S. (1978). Mind in Society: The Development of Higher Psychological Processes. Harvard University Press.",
        "Wood, D., Bruner, J. S., & Ross, G. (1976). The role of tutoring in problem solving. Journal of Child Psychology and Psychiatry, 17(2), 89–100.",
        "Collins, A., Brown, J. S., & Newman, S. E. (1989). Cognitive apprenticeship: Teaching the crafts of reading, writing, and mathematics. In L. B. Resnick (Ed.), Knowing, Learning, and Instruction (pp. 453–494). Erlbaum.",
        "Palincsar, A. S., & Brown, A. L. (1984). Reciprocal teaching of comprehension-fostering and comprehension-monitoring activities. Cognition and Instruction, 1(2), 117–175.",
        "Mercer, N. (1996). The quality of talk in children's collaborative activity in the classroom. Learning and Instruction, 6(4), 359–377.",
        "Osborne, J. (2010). Arguing to learn in science: The role of collaborative, critical discourse. Science, 328(5977), 463–466.",
        "Kuhn, D. (1991). The Skills of Argument. Cambridge University Press.",
        "Scardamalia, M., & Bereiter, C. (1994). Computer support for knowledge-building communities. Journal of the Learning Sciences, 3(3), 265–283.",
        "van Eemeren, F. H., & Grootendorst, R. (2004). A Systematic Theory of Argumentation: The Pragma-Dialectical Approach. Cambridge University Press.",
        "van Zee, E. H., & Minstrell, J. (1997). Using questioning to guide student thinking. Journal of the Learning Sciences, 6(2), 227–269.",
        "Paul, R., & Elder, L. (2007). Critical Thinking: Learn the Tools the Best Thinkers Use. Pearson.",
        "Graesser, A. C., & Person, N. K. (1994). Question asking during tutoring. American Educational Research Journal, 31(1), 104–137.",
        "Alfieri, L., Nokes-Malach, T. J., & Schunn, C. D. (2013). Learning through case comparisons: A meta-analytic review. Educational Psychologist, 48(2), 87–113.",
        "Chi, M. T. H. (2009). Active-constructive-interactive: A conceptual framework for differentiating learning activities. Topics in Cognitive Science, 1(1), 73–105.",
        "von Glasersfeld, E. (1989). Cognition, construction of knowledge, and teaching. Synthese, 80(1), 121–140.",
        "Facione, P. A. (1990). Critical Thinking: A Statement of Expert Consensus (Delphi Report). ERIC ED315423.",
        "Toulmin, S. E. (1958). The Uses of Argument. Cambridge University Press.",
        "Biggs, J. B., & Collis, K. F. (1982). Evaluating the Quality of Learning: The SOLO Taxonomy. Academic Press.",
        "Novak, J. D., & Gowin, D. B. (1984). Learning How to Learn. Cambridge University Press.",
        "Mislevy, R. J., Steinberg, L. S., & Almond, R. G. (2003). On the structure of educational assessments. Measurement, 1(1), 3–62.",
        "Brookhart, S. M. (2013). How to Create and Use Rubrics for Formative Assessment and Grading. ASCD.",
    ]:
        add_ref(doc, ref)

    out = "/Users/Zhuanz/Documents/Codex/2026-08-02/xia/outputs/教育理论基础文献清单.docx"
    doc.save(out)
    print("saved:", out)


def add_h2(doc, text):
    return add_heading(doc, text, 2)


if __name__ == "__main__":
    build()
